"""培养方案 docx 解析（仅安全工程 / 海洋油气工程 2 专业）。
产出：
  - fact_plan_course：计划课程明细（模块/学分/学期/核心课）
  - fact_plan_meta：专业级修读要求（总/必/选/实践学分 + 12 条毕业要求）
docx 结构（数据探明）：
  - 段落"修读要求"含 必修课/选修课/实践/最低总学分（Tab 分隔）
  - 段落"毕业生应获得的知识和能力"含编号 1..12
  - 1 张计划课程表：18 列，前 4 列为模块层级，数据从第 3 行起
"""
import json
import re
import docx
import pandas as pd

from . import config

# 计划课程表列索引（0 基）
_COL_MODULE = 1      # 课程模块（通识教育课/学科基础课/...）
_COL_CODE = 4
_COL_NAME = 5
_COL_NATURE = 6
_COL_CREDITS = 7
_COL_TERM = 14
_COL_REQUIRED = 15
_COL_DEPT = 16
_COL_NOTE = 17


def _num(s):
    m = re.search(r"-?\d+(\.\d+)?", str(s or ""))
    return float(m.group()) if m else None


def _find_para(doc, keyword):
    for p in doc.paragraphs:
        if keyword in p.text:
            return p.text
    return ""


def _find_body_after(doc, heading_kw, stop_kws=("计划课程", "五 、", "★")):
    """返回标题段之后、下个章节之前的正文（多段拼接）。"""
    paras = doc.paragraphs
    start = None
    for i, p in enumerate(paras):
        if heading_kw in p.text:
            start = i
            break
    if start is None:
        return ""
    body = []
    for p in paras[start + 1:]:
        if any(k in p.text for k in stop_kws):
            break
        if p.text.strip():
            body.append(p.text)
    return "\n".join(body)


def _core_course_names(doc):
    """从'主要课程'段提取核心课程名集合（顿号分隔）。"""
    txt = _find_body_after(doc, "主要课程", stop_kws=("毕业生应获得", "四 、"))
    names = set()
    for nm in re.split(r"[、，,；;\n]", txt):
        nm = nm.strip().strip("等。").strip()
        if nm and len(nm) >= 2:
            names.add(nm)
    return names


def parse_requirements(doc):
    """从'修读要求'段提取 总/必/选/实践 学分 + 学位要求。"""
    txt = _find_para(doc, "修读要求") + "\n" + _find_para(doc, "最低总学分")
    def grab(label):
        m = re.search(label + r"[^\d]*?(\d+(?:\.\d+)?)", txt)
        return float(m.group(1)) if m else None
    degree = ""
    m = re.search(r"学位[^\n]*?要求[^\n]*", txt)
    if m:
        degree = m.group().strip()
    return {
        "required_credits": grab("必修课"),
        "elective_credits": grab("选修课"),
        "practice_credits": grab("实践教学环节"),
        "total_credits": grab("最低总学分"),
        "degree_req": degree or "满足学校规定的学位授予条件",
    }


def parse_grad_requirements(doc):
    """从'毕业生应获得'段后的正文提取 12 条（编号 1. 2. 切分）。"""
    txt = _find_body_after(doc, "毕业生应获得", stop_kws=("计划课程", "五 、", "★表示"))
    if not txt:
        return []
    items = re.split(r"(?:^|\n)\s*(\d{1,2})[\.、]\s*", txt)
    reqs = []
    for i in range(1, len(items) - 1, 2):
        idx = items[i]
        body = items[i + 1].strip().replace("\n", " ")
        name = body.split("：")[0].split(":")[0][:20] if body else ""
        reqs.append({"id": int(idx), "name": name, "desc": body})
    return reqs


def parse_plan_courses(doc):
    """解析计划课程表 → list[dict]。过滤分节/合计等无课程代码的行。"""
    if not doc.tables:
        return []
    core_names = _core_course_names(doc)
    t = doc.tables[0]
    rows = []
    for ri in range(2, len(t.rows)):       # 前 2 行表头
        cells = [c.text.strip() for c in t.rows[ri].cells]
        if len(cells) <= _COL_DEPT:
            continue
        code = cells[_COL_CODE]
        name = cells[_COL_NAME]
        # 仅保留有合法课程代码的行（过滤"要求学分:166"等分节行）
        if not re.match(r"^[0-9A-Za-z]{6,}$", code or ""):
            continue
        clean_name = name.replace("★", "").replace("▲", "").strip()
        note = cells[_COL_NOTE] if len(cells) > _COL_NOTE else ""
        is_core = 1 if ("★" in name or "★" in note or clean_name in core_names) else 0
        rows.append({
            "course_id": code,
            "course_name": clean_name or None,
            "module": cells[_COL_MODULE] or None,
            "credits": _num(cells[_COL_CREDITS]),
            "term": cells[_COL_TERM] or None,
            "is_core": is_core,
            "source": "real",
        })
    return rows


def parse_plan(major_name: str):
    """解析单个专业培养方案 → (meta_dict, courses_list)。"""
    path = config.PLAN_DOCX[major_name]
    doc = docx.Document(str(path))
    req = parse_requirements(doc)
    reqs12 = parse_grad_requirements(doc)
    courses = parse_plan_courses(doc)
    meta = {
        "major_name": major_name,
        "grade": "2022",
        **req,
        "grad_reqs_json": json.dumps(reqs12, ensure_ascii=False),
        "source": "real",
    }
    return meta, courses


def parse_all():
    """解析 2 专业 → (meta_df, courses_df)（major_id 待 run_etl 映射）。"""
    metas, all_courses = [], []
    for name in config.PLAN_DOCX:
        meta, courses = parse_plan(name)
        metas.append(meta)
        for c in courses:
            c["major_name"] = name
            c["grade"] = "2022"
        all_courses.extend(courses)
    return pd.DataFrame(metas), pd.DataFrame(all_courses)


if __name__ == "__main__":
    mdf, cdf = parse_all()
    for _, m in mdf.iterrows():
        reqs = json.loads(m["grad_reqs_json"])
        print(f"\n【{m['major_name']}】总{m['total_credits']} 必{m['required_credits']} "
              f"选{m['elective_credits']} 实践{m['practice_credits']}  毕业要求{len(reqs)}条")
        print("  计划课程:", len(cdf[cdf['major_name'] == m['major_name']]), "门",
              "核心课:", int(cdf[cdf['major_name'] == m['major_name']]['is_core'].sum()))
    print("\n模块分布(安全工程):")
    print(cdf[cdf['major_name'] == '安全工程']['module'].value_counts().to_dict())
