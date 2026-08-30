"""基础报表 Excel 导出；内容与页面使用同一查询快照。"""
from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .catalog import REPORTS


LABELS = {
    "group": "分组", "organizationName": "学院", "majorCode": "专业代码", "majorName": "专业",
    "classCode": "班级", "studentCount": "学生人数", "validGradeStudents": "有效成绩学生数",
    "failedBeforeStudents": "补考前挂科人数", "failedBeforeRate": "补考前挂科率",
    "failedAfterStudents": "补考后仍挂科人数", "failedAfterRate": "补考后挂科率",
    "gradeCoverageRate": "成绩覆盖率", "retainedDemotedStudents": "留/降级人数", "makeupAttempts": "补考人次", "makeupPass": "补考通过人次",
    "makeupPassRate": "通过率", "failedStudents": "挂科人数", "failureRate": "挂科率",
    "studentCountDisplay": "专业人数", "failedBeforeDisplay": "已挂人数", "failedAfterDisplay": "在挂人数", "passRate": "通过率",
    "maleValidStudents": "男生有效成绩人数", "maleFailedStudents": "男生挂科人数", "maleFailureRate": "男生挂科率",
    "femaleValidStudents": "女生有效成绩人数", "femaleFailedStudents": "女生挂科人数", "femaleFailureRate": "女生挂科率",
    "unknownGenderStudents": "性别未说明人数", "zeroCourse": "0门", "oneToTwo": "1-2门",
    "threeToFive": "3-5门", "sixOrMore": "6门及以上", "rankedStudents": "有成绩排名人数",
    "top20": "专业前20%", "top20Rate": "专业前20%占比", "top20To50": "专业前20%-50%",
    "top20To50Rate": "专业前20%-50%占比", "top50To80": "专业前50%-80%",
    "top50To80Rate": "专业前50%-80%占比", "bottom20": "专业后20%", "bottom20Rate": "专业后20%占比",
    "noGradeStudents": "无有效成绩人数", "courseId": "课程代码", "courseName": "课程名称",
    "firstValidStudents": "首次有效结果人数", "cet4PassedStudents": "四级累计通过人数",
    "cet4PassRate": "四级累计通过率", "cet4PassRateRank": "四级通过率排行",
    "sequence": "序号", "studentId": "学号", "name": "姓名", "mentor": "导师", "ruleId": "规则",
    "level": "级别", "unresolvedCourseCount": "未解决课程门数", "courseEvidence": "课程证据", "triggerDetail": "触发说明",
}
RATE_KEYS = {"rate", "failedBeforeRate", "failedAfterRate", "failureRate", "makeupPassRate", "passRate",
             "failedBeforeCourseRate", "failedAfterCourseRate", "cet4PassRate"}

RANK_GLYPHS = {1: "①", 2: "②", 3: "③"}
RANK_COLORS = {1: "C9A227", 2: "B87333", 3: "9CA3AF"}


def _export_value(row: dict[str, Any], key: str) -> Any:
    if key == "cet4PassRateRank":
        return RANK_GLYPHS.get(row.get(key), row.get(key))
    return row.get(key)


def _style_header(ws, header_rows: int, start_row: int = 1) -> None:
    for row in ws.iter_rows(min_row=start_row, max_row=start_row + header_rows - 1):
        for cell in row:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="315D7C")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _merge_result_cells(ws, rows: list[dict[str, Any]], keys: list[str], merge_by: tuple[tuple[str, str], ...], data_start: int) -> None:
    key_columns = {key: index + 1 for index, key in enumerate(keys)}
    for visible_key, group_key in merge_by:
        column = key_columns.get(visible_key)
        if column is None:
            continue
        index = 0
        while index < len(rows):
            row = rows[index]
            group_value = row.get(group_key)
            if group_value in {None, ""} or row.get("isSummary"):
                index += 1
                continue
            end = index + 1
            while end < len(rows) and rows[end].get(group_key) == group_value and not rows[end].get("isSummary"):
                end += 1
            if end - index > 1:
                ws.merge_cells(start_row=data_start + index, start_column=column,
                               end_row=data_start + end - 1, end_column=column)
            index = end


def _write_table(ws, payload: dict) -> None:
    report = REPORTS[payload["reportId"]]
    rows = payload.get("rows") or []
    columns = list(report.result_columns)
    keys = [key for key, _ in columns]
    title_rows = 1 if report.report_id in {"RPT-01", "RPT-02", "RPT-03", "RPT-04A", "RPT-04B", "RPT-05", "RPT-06", "RPT-07", "RPT-08"} else 0
    header_rows = 2 if report.report_id == "RPT-05" else 1
    header_start = title_rows + 1
    if title_rows:
        ws.append([payload.get("title") or report.title] + [None] * (len(columns) - 1))
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
        title_cell = ws.cell(1, 1)
        title_cell.font = Font(bold=True, size=14)
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28
    if report.report_id == "RPT-05":
        ws.append(["科目", "专业", "专业人数", "补考前数据", None, None, None,
                   "补考后数据", None, None, None, "补考通过人数"])
        ws.append([None, None, None, "挂科人数", "专业挂科率", "挂科总人数", "总挂科率",
                   "挂科人数", "专业挂科率", "挂科总人数", "总挂科率", None])
        for column in (1, 2, 3, 12):
            ws.merge_cells(start_row=header_start, start_column=column,
                           end_row=header_start + 1, end_column=column)
        ws.merge_cells(start_row=header_start, start_column=4,
                       end_row=header_start, end_column=7)
        ws.merge_cells(start_row=header_start, start_column=8,
                       end_row=header_start, end_column=11)
    else:
        ws.append([label for _, label in columns])
    for row in rows:
        ws.append([_export_value(row, key) for key in keys])
    _style_header(ws, header_rows, header_start)
    data_start = title_rows + header_rows + 1
    rank_column = keys.index("cet4PassRateRank") + 1 if "cet4PassRateRank" in keys else None
    rank_glyphs = set(RANK_GLYPHS)
    for row_index, row in enumerate(rows, data_start):
        if row.get("isSummary"):
            for cell in ws[row_index]:
                cell.font = Font(bold=True)
        for cell in ws[row_index]:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        rank = row.get("cet4PassRateRank")
        if rank_column and rank in rank_glyphs and not row.get("isSummary"):
            rank_cell = ws.cell(row_index, rank_column)
            rank_cell.font = Font(bold=True, size=14, color=RANK_COLORS[rank])
            rank_cell.alignment = Alignment(horizontal="center", vertical="center")
    for col_index, (key, label) in enumerate(columns, 1):
        letter = get_column_letter(col_index)
        ws.column_dimensions[letter].width = min(48, max(12, len(label) * 2 + 2))
        if key in RATE_KEYS:
            for row_index in range(data_start, ws.max_row + 1):
                ws.cell(row_index, col_index).number_format = "0.00%"
        if key in {"studentId", "classCode"}:
            for row_index in range(data_start, ws.max_row + 1):
                ws.cell(row_index, col_index).number_format = "@"
    _merge_result_cells(ws, rows, keys, report.merge_by, data_start)
    ws.freeze_panes = f"A{data_start}"
    if rows and report.report_id != "RPT-05":
        ws.auto_filter.ref = f"A{header_start}:{ws.cell(ws.max_row, len(keys)).coordinate}"


def build_workbook(payload: dict, user: dict) -> bytes:
    workbook = Workbook()
    result = workbook.active
    result.title = "报表结果"
    _write_table(result, payload)

    explanation = workbook.create_sheet("报表说明")
    context = payload.get("context") or {}
    entry_grade = context.get("entryGrade")
    grade_value = entry_grade
    if entry_grade is None and payload.get("reportId") == "RPT-05":
        grade_value = "全部年级"
    elif entry_grade is None and payload.get("reportId") in {"RPT-07", "RPT-08"}:
        grade_value = "全部授权年级"
    explanation_rows = [
        ("报表编号", payload.get("reportId")), ("报表标题", payload.get("title")),
        ("可用状态", payload.get("status")), ("导出账号", user.get("username")),
        ("当前工作身份", context.get("identity")), ("学年学期", context.get("semesterId")),
        ("年级" if payload.get("reportId") in {"RPT-02", "RPT-03", "RPT-04A", "RPT-04B", "RPT-05", "RPT-06", "RPT-07", "RPT-08"} else "入学年级", grade_value),
        ("学院筛选", context.get("organizationId") or "全部授权范围"),
        ("专业筛选", context.get("majorCode") or "全部授权范围"), ("班级筛选", context.get("classCode") or "全部授权范围"),
        ("口径版本", context.get("ruleVersion")), ("数据截止时间", str(context.get("dataCutoff"))),
        ("边界说明", "；".join(payload.get("boundary") or [])),
    ]
    focus_rule = payload.get("focusRule") or {}
    if payload.get("reportId") in {"RPT-07", "RPT-08"}:
        descriptions = [
            f"{item.get('ruleId')}（{item.get('level')}）：{item.get('description')}"
            f"{'（当前启用）' if item.get('enabled') else '（当前停用）'}"
            for item in focus_rule.get("items") or []
        ]
        explanation_rows.extend([
            ("重点关注规则", "；".join(descriptions)),
            ("预警规则版本", "、".join(focus_rule.get("versions") or []) or "未标注"),
            ("课程去重口径", focus_rule.get("dedupDescription")),
        ])
    for label, value in explanation_rows:
        explanation.append([label, value])
    explanation.column_dimensions["A"].width = 22
    explanation.column_dimensions["B"].width = 100
    for cell in explanation["A"]:
        cell.font = Font(bold=True)

    rules = workbook.create_sheet("口径规则")
    rules.append(["规则编号", "来源类型", "来源", "公式", "适用边界"])
    for rule in payload.get("rules") or []:
        rules.append([rule.get("ruleId"), rule.get("provenance"), rule.get("source"), rule.get("formula"), rule.get("boundary")])
    for cell in rules[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="315D7C")
    for column, width in zip("ABCDE", (22, 16, 40, 70, 70)):
        rules.column_dimensions[column].width = width
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
