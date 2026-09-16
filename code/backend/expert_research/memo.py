"""One immutable reading structure shared by web preview and editable export."""
from copy import deepcopy
from datetime import datetime, timezone, timedelta
from .program_review import disputed


def local_time(value):
    if not value: return '未提供'
    try:
        stamp=datetime.fromisoformat(value.replace('Z','+00:00'))
        if stamp.tzinfo is None: return value
        return stamp.astimezone(timezone(timedelta(hours=8))).strftime('%Y年%m月%d日 %H:%M')
    except (TypeError, ValueError): return str(value)


def build(snapshot):
    r=snapshot['result']; sections=[]
    def section(title, lines):
        lines=[s for s in lines if s]
        if lines: sections.append({'title':title,'paragraphs':lines})
    bad={d['plan'] for d in r.get('documents',[]) if disputed(d)}
    issues=[d['plan']+'：'+'；'.join(d.get('issues',[])) for d in r.get('documents',[]) if disputed(d)]
    scope=r.get('scope') or {}
    plan_by_id={p['plan_id']:p for p in snapshot.get('frozen_plans',[]) if p.get('plan_id')}
    labels=[]
    for key in ('plan_id','target_plan_id'):
        p=plan_by_id.get(scope.get(key))
        if not p: continue
        label=p.get('plan_name') or p.get('major_name') or '名称未提供'
        grade=str(p.get('grade') or '')
        if grade and grade not in label:label+=' · '+grade+'级'
        labels.append(label)
    original=snapshot.get('research_question') or snapshot.get('question') or snapshot.get('title','研究事项')
    current=snapshot.get('question')
    section('研究事项',[original, '本轮关注：'+current if current and current!=original else '',
        '研究对象：'+' 与 '.join(labels) if labels else '',
        '成绩学期：'+scope['semester'] if scope.get('semester') else '',
        r.get('request_receipt',''),
        r.get('scope_label','当前授权范围')+'；资料导入：'+local_time(r.get('data_time'))+'。业务截至日期尚未确认。'])
    section('目前看法', [('培养目标原文归属有冲突，暂不据此判断专业特色；课程结构只作限定对照。' if issues else r.get('headline','')),
        r.get('decision_summary'),r.get('body')])
    opinion=snapshot.get('opinion') or {}
    section('个人意见（经本人确认选入）',[opinion.get('text'), f"采用已保存意见版本 {opinion.get('revision',0)}。" if opinion.get('text') else ''])
    points=r.get('discussion_points',[])
    section('建议讨论的事项',[p['title']+'：'+p['detail']+'\n需了解：'+p.get('needed','') for p in points] or [r.get('management_note')])
    for t in r.get('tables',[]):
        if t['id'] not in {'course_options','course_current','transition_summary','bottlenecks','readiness'}: continue
        lines=['；'.join(col['label']+'：'+str(row.get(col['key']) if row.get(col['key']) is not None else '未提供') for col in t['columns']) for row in t['rows'][:3]]
        if len(t['rows'])>3:lines.append('按原表顺序列出前3项，完整情况见附录；此处不代表全部问题。')
        if t.get('note'):lines.append(t['note'])
        section(t['title']+'（摘要）',lines)
    for c in r.get('contributions',[]):
        other=c['result']; section(other.get('title','补充分析'),[other.get('headline'),other.get('body')])
    questions=[q['text']+('（暂缓）' if q.get('status')=='deferred' else '') for q in snapshot.get('questions',[]) if q.get('status')!='resolved']
    section('尚需补充的情况',list(dict.fromkeys(issues+questions+r.get('missing',[]))))
    section('使用条件',list(dict.fromkeys(r.get('limitations',[])+['本稿仅供研究讨论，不是学校审批、课程认定或资格决定。'])))
    tables=deepcopy(r.get('tables',[]))
    for t in tables:
        if t['id'] in {'positioning','named_courses','reading_requirements'}:
            t['rows']=[row for row in t['rows'] if row.get('plan') not in bad]
    title={'program':'专业培养方案对照','course':'课程质量建设研究','transfer':'转专业培养衔接研究',
           'graduation':'毕业准备情况研究','recommendation':'推免工作准备研究'}.get(r.get('expert_id'),'教学管理研究')
    return {'version':'discussion-memo/2.2','title':title,'sections':sections,'appendix_tables':tables,
            'methods':r.get('methods',[]),'sources':[{'file':d['file'],'detail':d['plan']+('；归属待确认，不用于定位判断' if d['plan'] in bad else '')} for d in r.get('documents',[])]+r.get('sources',[])}


def render(material):
    from io import BytesIO
    from docx import Document
    from docx.shared import Mm, Pt
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    doc=Document(); page=doc.sections[0]
    page.page_width=Mm(210); page.page_height=Mm(297)
    page.left_margin=page.right_margin=Mm(24);page.top_margin=page.bottom_margin=Mm(22)
    for name,size in [('Normal',11),('Title',18),('Heading 1',14),('Heading 2',12)]:
        style=doc.styles[name];style.font.name='Microsoft YaHei';style.font.size=Pt(size)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    doc.styles['Normal'].paragraph_format.line_spacing=1.5
    memo=material['snapshot']['discussion_memo']
    doc.add_paragraph((memo.get('title') or material.get('title','研究事项'))+' · 讨论稿','Title')
    for section in memo['sections']:
        doc.add_heading(section['title'],level=1)
        for text in section['paragraphs']: doc.add_paragraph(text)
    doc.add_page_break();doc.add_heading('附录：课程与计算资料',level=1)
    for item in memo['appendix_tables']:
        if not item['rows']:continue
        doc.add_heading(item['title'],level=2)
        cols=item['columns'][:4];table=doc.add_table(rows=1,cols=len(cols));table.style='Table Grid'
        for i,col in enumerate(cols):table.rows[0].cells[i].text=col['label']
        repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for record in item['rows'][:15]:
            cells=table.add_row().cells
            for i,col in enumerate(cols):cells[i].text=str(record.get(col['key']) if record.get(col['key']) is not None else '未提供')
        if len(item['rows'])>15 or len(item['columns'])>4:doc.add_paragraph('此处为摘要（最多15行、4列），完整字段和清单保存在本材料网页附录及对应研究中。')
        if item.get('note'):doc.add_paragraph(item['note'])
    doc.add_heading('来源与计算说明',level=1)
    for source in memo['sources']:doc.add_paragraph(source['file']+'：'+source.get('detail',''))
    for text in memo['methods']:doc.add_paragraph(text)
    doc.add_paragraph('形成时间：'+local_time(material.get('created_at')))
    doc.core_properties.identifier=material['id']
    doc.core_properties.author='';doc.core_properties.last_modified_by=''
    output=BytesIO();doc.save(output);return output.getvalue()
