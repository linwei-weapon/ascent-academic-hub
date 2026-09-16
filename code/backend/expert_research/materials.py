"""Editable discussion material from one immutable, authorized snapshot."""
from io import BytesIO


def render_docx(material):
    if material['snapshot'].get('discussion_memo'):
        from .memo import render
        return render(material)
    from docx import Document
    from docx.shared import Mm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    snap=material['snapshot'];result=snap['result']
    doc=Document();section=doc.sections[0]
    section.page_width=Mm(210);section.page_height=Mm(297)
    section.top_margin=section.bottom_margin=Mm(24)
    section.left_margin=section.right_margin=Mm(24)
    for name,size in [('Normal',12),('Title',18),('Heading 1',14),('Heading 2',12)]:
        style=doc.styles[name];style.font.name='Microsoft YaHei';style.font.size=Pt(size)
        style.font.color.rgb=RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    normal=doc.styles['Normal'].paragraph_format
    normal.line_spacing=1.5;normal.space_after=Pt(6)
    doc.add_paragraph((material.get('title') or result.get('title') or '研究事项')+'讨论稿','Title')
    scope=result.get('scope',{})
    doc.add_paragraph(f"范围：{result.get('scope_label','当前授权范围')}。适用方案：{(result.get('comparison') or {}).get('source',{}).get('plan_name') or scope.get('plan_id','未提供')}。")
    doc.add_paragraph(f"学期：{scope.get('semester') or '未指定'}。业务截至：未提供。资料导入时间：{result.get('data_time') or '未提供'}。")
    doc.add_paragraph('本稿供研究和讨论，依据已接入资料形成，不是学校审批或资格认定。')
    doc.add_heading('目前可以明确的情况',level=1)
    for key in ('headline','body','management_note'):
        if result.get(key):doc.add_paragraph(result[key])
    for c in result.get('contributions',[]):
        r=c['result'];doc.add_heading(r['title'],level=2)
        doc.add_paragraph(r['headline']);doc.add_paragraph(r.get('body',''))
    opinion=snap.get('opinion') or {}
    if opinion.get('text'):
        doc.add_heading('当前想法',level=1)
        for line in opinion['text'].splitlines():doc.add_paragraph(line)
    questions=snap.get('questions') or []
    if questions:
        doc.add_heading('尚需明确的事项',level=1)
        for i,q in enumerate(questions,1):
            doc.add_paragraph(f"{i}. {q.get('text','')}"+('（暂缓）' if q.get('status')=='deferred' else ''))
    boundaries=list(dict.fromkeys(result.get('missing',[])+result.get('limitations',[])))
    if boundaries:
        doc.add_heading('使用条件',level=1)
        for text in boundaries:doc.add_paragraph(text)
    # Deliberately small comparable tables. Wide and long course details stay
    # in the saved research rather than shrinking the memo into unreadable type.
    selected=[t for t in result.get('tables',[]) if t['id'] in {'comparison','layers','course_current','readiness','transition_summary','graduation_changes'}]
    for source in selected[:3]:
        doc.add_heading(source['title'],level=1)
        columns=source['columns'][:4]
        table=doc.add_table(rows=1,cols=len(columns));table.alignment=WD_TABLE_ALIGNMENT.CENTER
        table.autofit=False
        for j,c in enumerate(columns):
            table.columns[j].width=Mm(162/len(columns));table.rows[0].cells[j].text=c['label']
        header=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(header)
        for record in source['rows'][:15]:
            row=table.add_row()
            for j,c in enumerate(columns):
                value=record.get(c['key']);row.cells[j].text='未提供' if value is None else str(value)
        for i,row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                props=cell._tc.get_or_add_tcPr()
                borders=OxmlElement('w:tcBorders')
                for edge in ('top','left','bottom','right'):
                    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
                props.append(borders)
                margins=OxmlElement('w:tcMar')
                for edge in ('top','left','bottom','right'):
                    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'100');el.set(qn('w:type'),'dxa');margins.append(el)
                props.append(margins)
                if i==0:
                    fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E8EDF2');props.append(fill)
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing=1.2;p.paragraph_format.space_after=Pt(3)
                    for run in p.runs:run.font.size=Pt(10);run.bold=(i==0)
        if len(source['rows'])>15 or len(source['columns'])>4:
            doc.add_paragraph('此处列出摘要，完整明细保存在对应研究结果中。')
        if source.get('note'):doc.add_paragraph(source['note'])
    doc.add_heading('计算口径与来源',level=1)
    for m in result.get('methods',[]):doc.add_paragraph(m)
    for source in result.get('documents',[]):
        doc.add_paragraph(f"{source.get('file','资料')}；适用方案：{source.get('plan','未提供')}；保存版本：{source.get('hash') or '未提供'}。")
    doc.add_paragraph(f"研究编号：{material.get('research_id','')}\n材料版本：{material['id']}\n形成时间：{material.get('created_at','')}\n分析版本：{result.get('version','')}\n意见版本：{opinion.get('revision',0)}")
    doc.core_properties.author='';doc.core_properties.last_modified_by=''
    output=BytesIO();doc.save(output);return output.getvalue()
